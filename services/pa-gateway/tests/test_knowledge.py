# -*- coding: utf-8 -*-
"""WP-07: the knowledge pipeline's real behaviour, without touching the index."""
import pytest

import app.knowledge as knowledge_mod
from app.db import DbStore
from app.knowledge import (KnowledgeError, KnowledgeService, chunk_text,
                           extract_text)


@pytest.fixture()
def service(tmp_path, monkeypatch):
    store = DbStore(path=str(tmp_path / "knowledge.db"))
    indexed: list = []
    monkeypatch.setattr(knowledge_mod, "index_chunks",
                        lambda document, chunks: (indexed.append((document, chunks)),
                                                  len(chunks))[1])
    monkeypatch.setattr(knowledge_mod, "remove_document",
                        lambda document_id: {"deleted": 2})
    monkeypatch.setattr(knowledge_mod, "drop_collection", lambda: True)
    monkeypatch.setattr(knowledge_mod, "ensure_collection", lambda: {"created": True})
    return KnowledgeService(store), store, indexed, tmp_path


def _write(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


# ------------------------------------------------------------------- parsing


def test_chunks_never_span_two_pages():
    text = "a" * 1200
    chunks = chunk_text(text, sections=[(1, "a" * 100), (2, "b" * 100)], size=60,
                        overlap=10)
    pages = {chunk["page"] for chunk in chunks}
    assert pages == {1, 2}
    assert all(chunk["page"] in (1, 2) for chunk in chunks)
    assert chunks[0]["text"].startswith("a")
    assert chunks[-1]["text"].endswith("b")


def test_chunks_overlap_so_a_sentence_is_not_cut_in_half():
    text = "".join(str(i % 10) for i in range(200))
    chunks = chunk_text(text, size=100, overlap=20)
    assert len(chunks) >= 2
    assert chunks[0]["text"][-20:] == chunks[1]["text"][:20]


def test_extract_text_reads_a_real_markdown_file(tmp_path):
    path = _write(tmp_path, "note.md", "# Elfred\n北极星目标：让 PA 代表用户做事")
    text, sections = extract_text(path)
    assert "北极星目标" in text and sections == []


def test_extract_text_refuses_an_unsupported_format(tmp_path):
    path = tmp_path / "raw.bin"
    path.write_bytes(b"\x00\x01\x02")
    with pytest.raises(KnowledgeError) as exc:
        extract_text(path)
    assert exc.value.code == "UNSUPPORTED_FORMAT"


def test_extract_text_reports_a_missing_file(tmp_path):
    with pytest.raises(KnowledgeError) as exc:
        extract_text(tmp_path / "nope.md")
    assert exc.value.code == "FILE_NOT_FOUND"


def test_extract_text_parses_a_real_pptx(tmp_path):
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Elfred 北极星目标"
    path = tmp_path / "deck.pptx"
    deck.save(str(path))
    text, sections = extract_text(path)
    assert "北极星目标" in text
    assert sections and sections[0][0] == 1


# ------------------------------------------------------------------ lifecycle


def test_ingest_parses_chunks_stores_and_indexes(service):
    svc, store, indexed, tmp_path = service
    path = _write(tmp_path, "plan.md", "北极星目标：让每个用户的 PA 代表用户做事。\n"
                                      "第二步是从 PA 走到 A2A 交易平台。")
    result = svc.ingest(str(path), title="Phase-1 蓝图", kind="personal", pa_id="pa_x")
    document = result["document"]
    assert document["status"] == "indexed"
    assert document["chunk_count"] == result["chunks"] >= 1
    assert document["checksum"] and document["bytes"] > 0
    chunks = store.list_chunks(document["document_id"])
    assert chunks and chunks[0]["text"].startswith("北极星目标")
    assert indexed[0][0]["pa_id"] == "pa_x"
    jobs = store.list_jobs()
    assert jobs[0]["state"] == "succeeded" and jobs[0]["kind"] == "ingest"


def test_a_failed_index_keeps_the_chunks_and_says_so(service, monkeypatch):
    svc, store, _indexed, tmp_path = service

    def boom(document, chunks):
        raise KnowledgeError("INDEX_UNAVAILABLE", "typesense is down")

    monkeypatch.setattr(knowledge_mod, "index_chunks", boom)
    path = _write(tmp_path, "plan.md", "北极星目标")
    with pytest.raises(KnowledgeError):
        svc.ingest(str(path), pa_id="pa_x")
    document = store.list_documents()[0]
    assert document["status"] == "stored"
    assert "INDEX_UNAVAILABLE" in document["error"]
    assert store.list_chunks(document["document_id"]), "the parsed chunks are kept"
    assert store.list_jobs()[0]["state"] == "failed"


def test_an_empty_document_fails_instead_of_indexing_nothing(service):
    svc, store, _indexed, tmp_path = service
    path = _write(tmp_path, "empty.md", "   \n  \n")
    with pytest.raises(KnowledgeError) as exc:
        svc.ingest(str(path), pa_id="pa_x")
    assert exc.value.code == "EMPTY_DOCUMENT"
    assert store.list_documents()[0]["status"] == "failed"


def test_delete_removes_index_and_rows(service):
    svc, store, _indexed, tmp_path = service
    document_id = svc.ingest(str(_write(tmp_path, "plan.md", "北极星目标")),
                             pa_id="pa_x")["document"]["document_id"]
    result = svc.delete(document_id)
    assert result["indexed_deleted"] == 2
    assert store.get_document(document_id) is None
    assert store.list_chunks(document_id) == []


def test_delete_an_unknown_document_is_not_a_silent_success(service):
    svc, _store, _indexed, _tmp = service
    with pytest.raises(KnowledgeError) as exc:
        svc.delete("kdoc_missing")
    assert exc.value.code == "NOT_FOUND"


def test_rebuild_re_indexes_every_stored_chunk(service):
    svc, store, indexed, tmp_path = service
    svc.ingest(str(_write(tmp_path, "a.md", "北极星目标 A")), pa_id="pa_x")
    svc.ingest(str(_write(tmp_path, "b.md", "北极星目标 B")), pa_id="pa_x")
    indexed.clear()
    result = svc.rebuild()
    assert len(result["reindexed"]) == 2
    assert len(indexed) == 2
    assert store.list_jobs()[0]["kind"] == "rebuild"
    assert store.list_jobs()[0]["state"] == "succeeded"


def test_rebuild_skips_a_document_with_no_chunks(service):
    svc, store, _indexed, tmp_path = service
    source = store.create_source("personal", "personal", "pa_x", "elfred_knowledge")
    document = store.create_document(source_id=source["source_id"], title="orphan",
                                     path="x")
    store.update_document(document["document_id"], status="stored")
    result = svc.rebuild()
    assert result["reindexed"] == [] and len(result["skipped"]) == 1


# -------------------------------------------------------------------- search


def test_pa_search_covers_its_own_documents_and_the_capability_library(monkeypatch):
    captured = {}

    class _Response:
        status_code = 200
        text = "{}"

        @staticmethod
        def json():
            return {"found": 1, "hits": [{"document": {
                "chunk_id": "kdoc_1::c0", "document_id": "kdoc_1", "title": "Phase-1",
                "page": 2, "kind": "personal", "text": "北极星目标"}}]}

    def fake_request(method, path, **kw):
        captured["method"], captured["path"], captured["params"] = method, path, kw.get("params")
        return _Response()

    monkeypatch.setattr(knowledge_mod, "_request", fake_request)
    result = knowledge_mod.search("北极星", pa_id="pa_x")
    assert "pa_id:=pa_x" in captured["params"]["filter_by"]
    assert "kind:=capability" in captured["params"]["filter_by"]
    hit = result["hits"][0]
    assert hit["chunk_id"] == "kdoc_1::c0" and hit["page"] == 2
    assert hit["document"]["id"] == "kdoc_1"


def test_search_on_a_missing_collection_is_empty_not_an_error(monkeypatch):
    class _Response:
        status_code = 404
        text = "{}"

        @staticmethod
        def json():
            return {}

    monkeypatch.setattr(knowledge_mod, "_request", lambda *a, **kw: _Response())
    assert knowledge_mod.search("anything") == {
        "found": 0, "hits": [], "note": "knowledge collection does not exist yet"}
